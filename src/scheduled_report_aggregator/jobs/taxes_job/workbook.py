"""State filing templates: the bundled official copies, the drift check against the live ones, and row writing."""

# Standard library imports
from dataclasses import dataclass
from importlib.resources import files
from io import BytesIO
from typing import TYPE_CHECKING
from urllib.request import Request, urlopen

# Third party imports
from openpyxl import load_workbook

if TYPE_CHECKING:
  # Standard library imports
  from collections.abc import Iterable, Sequence
  from datetime import date
  from decimal import Decimal

  # Third party imports
  from openpyxl.workbook.workbook import Workbook
  from openpyxl.worksheet.worksheet import Worksheet


__all__ = ["CellValue", "TemplateSpec", "compare_templates", "load_template", "template_drift", "write_rows"]

type CellValue = str | int | float | Decimal | date | None

# michigan.gov answers 403 to clients that don't look like a browser (checked 2026-10-15).
BROWSER_USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36"

MAX_CELLS_REPORTED = 5  # per sheet, so one reworded header row doesn't bury the rest of the review notes


@dataclass(frozen=True)
class TemplateSpec:
  """One state-issued filing workbook: the bundled copy under `templates/<state>/` and where the state publishes it.

  `header_rows` is how many leading rows of each schedule sheet are template (titles, headers, guidance);
  the drift check compares exactly those. `ignored_sheets` are compared by name only: Michigan's
  April 2026 revision rewrapped its Instructions text without touching any schedule.
  """

  state: str
  filename: str
  live_url: str
  header_rows: int
  ignored_sheets: tuple[str, ...] = ()


def load_template(spec: TemplateSpec) -> Workbook:
  """A fresh in-memory copy of the bundled template; the packaged file itself is never written."""
  data = files("scheduled_report_aggregator.jobs.taxes_job").joinpath("templates", spec.state, spec.filename).read_bytes()
  return load_workbook(BytesIO(data))


def template_drift(spec: TemplateSpec) -> list[str]:
  """Differences between the bundled template and the state's live copy, as review warnings.

  Never raises: the check informs the reviewer and must not stop a filing, so a failed download
  becomes a warning too and the job carries on with the bundled copy.
  """
  try:
    with urlopen(Request(spec.live_url, headers={"User-Agent": BROWSER_USER_AGENT}), timeout=30) as response:
      live = load_workbook(BytesIO(response.read()))
  except Exception as e:  # noqa: BLE001 - see docstring
    return [f"{spec.filename}: couldn't fetch the state's live template to check for changes ({type(e).__name__}: {e})."]
  return compare_templates(spec, load_template(spec), live)


def compare_templates(spec: TemplateSpec, bundled: Workbook, live: Workbook) -> list[str]:
  """Sheet-name and template-row differences between two copies of one template, as warnings (empty when equal)."""
  name = spec.filename
  warnings = [f"{name}: the state's live template adds sheet {s!r}." for s in live.sheetnames if s not in bundled.sheetnames]
  warnings += [f"{name}: the state's live template no longer has sheet {s!r}." for s in bundled.sheetnames if s not in live.sheetnames]

  for sheet in bundled.sheetnames:
    if sheet in spec.ignored_sheets or sheet not in live.sheetnames:
      continue
    ours, theirs = bundled[sheet], live[sheet]
    width = max(ours.max_column, theirs.max_column)
    changed = [
      ours.cell(r, c).coordinate
      for r in range(1, spec.header_rows + 1)
      for c in range(1, width + 1)
      if ours.cell(r, c).value != theirs.cell(r, c).value
    ]
    if changed:
      shown = ", ".join(changed[:MAX_CELLS_REPORTED]) + (
        f" and {len(changed) - MAX_CELLS_REPORTED} more" if len(changed) > MAX_CELLS_REPORTED else ""
      )
      warnings.append(f"{name}: sheet {sheet} changed in the state's live template ({shown}). Review before filing.")
  return warnings


def write_rows(ws: Worksheet, header: Sequence[str], rows: Iterable[Sequence[CellValue]], number_formats: Sequence[str | None]) -> int:
  """Write *rows* below *ws*'s header row, starting at the first empty row after it; return how many were written.

  The header row is found by its text, not assumed: a layout change in a newer template then fails
  here, loudly, instead of filling the wrong cells. Michigan's uploader stops at the first blank row,
  so rows go in contiguously.
  """
  expected = [h.strip() for h in header]
  header_row = next(
    (r for r in range(1, ws.max_row + 1) if [str(ws.cell(r, c + 1).value or "").strip() for c in range(len(expected))] == expected),
    None,
  )
  if header_row is None:
    msg = f"sheet {ws.title}: no row reads {expected}; the template's layout has changed"
    raise ValueError(msg)

  first = header_row + 1
  while ws.cell(first, 1).value not in (None, ""):
    first += 1

  count = 0
  for count, row in enumerate(rows, start=1):
    for c, (value, fmt) in enumerate(zip(row, number_formats, strict=True), start=1):
      cell = ws.cell(first + count - 1, c, value)
      if fmt is not None:
        cell.number_format = fmt
  return count
